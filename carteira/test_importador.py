"""Testes do importador de Carteira via Excel (`carteira/importador.py`).

O relatório do ERP tem linhas de título variáveis antes do cabeçalho de
verdade e mistura tipos na mesma coluna (Cod. Prod. numérico e
alfanumérico) — os casos aqui replicam exatamente essas armadilhas, vistas
no arquivo real usado para desenhar o parser (`Carteira_de_pedidos_
11-09.xls`), não hipóteses."""
import io
from datetime import date, datetime

import openpyxl
from django.test import SimpleTestCase

from . import importador

_CABECALHO = [
    "Data Emissão", "Nota", "Pedido", "CFOP", "Destinatário", "Município",
    "Frete", "Vendedor", "Quantidade", "Valor Unit.", "Valor Total", "ST",
    "Cod. Prod.", "Venc. 1ª Fatura", "Descrição do Produto", "Centro de Custo",
]


# Sentinela pra distinguir "não passei o parâmetro" (usa a data padrão) de
# "passei None de propósito" (célula de data vazia de verdade) — usar
# `data or padrao` faria as duas coisas serem tratadas igual, já que None
# é falsy (foi exatamente o bug que este comentário documenta).
_DATA_PADRAO = object()


def _linha_padrao(pedido=7784, quantidade=10, valor_unit=8.2, valor_total=82.0,
                   cod_prod=18684, data=_DATA_PADRAO, cliente="VESTIS CONFECCOES LTDA",
                   municipio="AMERICANA-SP", descricao="CORTINA VENEZA 2,40X1,70 HAVAN"):
    if data is _DATA_PADRAO:
        data = datetime(2026, 8, 24)
    return [
        data, None, pedido, 5124, cliente, municipio,
        "NI", 1, quantidade, valor_unit, valor_total, 0, cod_prod,
        "12/30/1899", descricao, "GGTTEX",
    ]


def _montar_planilha(linhas_dados: list[list], titulo_linhas: int = 3) -> io.BytesIO:
    """Monta um .xlsx em memória: `titulo_linhas` linhas em branco/título
    antes do cabeçalho real, depois o cabeçalho, depois os dados — mesma
    forma do relatório de verdade."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(titulo_linhas):
        ws.append([])
    ws.append(_CABECALHO)
    for linha in linhas_dados:
        ws.append(linha)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


class LocalizarCabecalhoTests(SimpleTestCase):
    def test_encontra_com_numero_variavel_de_linhas_de_titulo(self):
        for titulo_linhas in (0, 1, 3, 4):
            with self.subTest(titulo_linhas=titulo_linhas):
                arquivo = _montar_planilha([_linha_padrao()], titulo_linhas=titulo_linhas)
                resultado = importador.parse_excel_carteira(arquivo)
                self.assertEqual(resultado.linhas_importadas, 1)
                self.assertEqual(resultado.avisos, [])

    def test_arquivo_sem_cabecalho_reconhecivel_da_erro_claro(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Isto", "não", "é", "o", "relatório"])
        ws.append([1, 2, 3, 4, 5])
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        with self.assertRaises(importador.CabecalhoNaoEncontrado):
            importador.parse_excel_carteira(buffer)


class ParseLinhaTests(SimpleTestCase):
    def test_cod_prod_alfanumerico_nao_quebra(self):
        arquivo = _montar_planilha([_linha_padrao(cod_prod="1.72613.01.9999")])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertEqual(resultado.df.iloc[0]["COD_PROD"], "1.72613.01.9999")

    def test_pedido_numerico_vira_string_sem_ponto_zero(self):
        # Uma segunda linha com PEDIDO em branco força a coluna inteira a
        # virar float (NaN só existe em float) — reproduz o mix real do
        # Excel (maioria int, mas a coluna vira float quando tem lacuna).
        arquivo = _montar_planilha([_linha_padrao(pedido=7784), _linha_padrao(pedido=None)])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertEqual(resultado.df.iloc[0]["PEDIDO"], "7784")
        self.assertNotIn(".0", resultado.df.iloc[0]["PEDIDO"])

    def test_linha_sem_pedido_vira_aviso_nao_excecao(self):
        arquivo = _montar_planilha([_linha_padrao(pedido=7784), _linha_padrao(pedido=None)])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertEqual(resultado.linhas_ignoradas, 1)
        self.assertIn("sem PEDIDO", resultado.avisos[0])

    def test_quantidade_e_valor_total_zerados_vira_aviso(self):
        arquivo = _montar_planilha([
            _linha_padrao(pedido=1),
            _linha_padrao(pedido=2, quantidade=0, valor_total=0),
        ])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertTrue(any("zerados" in a for a in resultado.avisos))

    def test_data_ausente_vira_aviso(self):
        arquivo = _montar_planilha([_linha_padrao(pedido=1), _linha_padrao(pedido=2, data=None)])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertTrue(any("sem data" in a for a in resultado.avisos))

    def test_linha_totalmente_vazia_nao_vira_aviso(self):
        """Comum no fim do relatório — ruído da planilha, não merece
        aparecer como pendência pro usuário resolver."""
        arquivo = _montar_planilha([_linha_padrao(pedido=1), [None] * 16])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertEqual(resultado.avisos, [])

    def test_coluna_desconhecida_vira_aviso_mas_nao_trava(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([])
        ws.append([*_CABECALHO, "Coluna Nova Do ERP"])
        ws.append([*_linha_padrao(), "algo"])
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        resultado = importador.parse_excel_carteira(buffer)
        self.assertEqual(resultado.linhas_importadas, 1)
        self.assertTrue(any("desconhecida" in a for a in resultado.avisos))

    def test_colunas_ignoradas_conhecidas_nao_viram_aviso(self):
        """CFOP, Frete, ST e Venc. 1ª Fatura sempre aparecem no relatório
        real sem informação útil — não podem virar ruído a cada import."""
        arquivo = _montar_planilha([_linha_padrao()])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.avisos, [])

    def test_nota_em_branco_nao_quebra(self):
        arquivo = _montar_planilha([_linha_padrao()])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.df.iloc[0]["NOTA"], "")

    def test_estado_derivado_do_municipio(self):
        arquivo = _montar_planilha([_linha_padrao(municipio="ITAQUAQUECETUBA-SP")])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.df.iloc[0]["ESTADO"], "SP")

    def test_categoria_derivada_da_descricao(self):
        arquivo = _montar_planilha([_linha_padrao(descricao="LENCOL QUEEN C/ELAS.LISO SORT.")])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertEqual(resultado.df.iloc[0]["CATEGORIA"], "LENÇOL")


class FormatoDeSaidaTests(SimpleTestCase):
    """O DataFrame produzido precisa ter as mesmas colunas derivadas que
    `servicos.carregar_carteira_do_sheets()` já produz — os dois caminhos
    alimentam a mesma tabela sincronizada, então o formato tem que bater."""

    def test_colunas_derivadas_presentes(self):
        arquivo = _montar_planilha([_linha_padrao(data=datetime(2026, 3, 15))])
        resultado = importador.parse_excel_carteira(arquivo)
        df = resultado.df
        for coluna in ("ANO", "MES", "ANO_MES", "MES_LABEL", "CLIENTE", "CLIENTE_CURTO"):
            self.assertIn(coluna, df.columns)
        self.assertEqual(df.iloc[0]["ANO"], 2026)
        self.assertEqual(df.iloc[0]["MES"], 3)
        self.assertEqual(df.iloc[0]["ANO_MES"], "2026-03")

    def test_sem_registros_validos_devolve_df_vazio_com_avisos(self):
        arquivo = _montar_planilha([_linha_padrao(pedido=None)])
        resultado = importador.parse_excel_carteira(arquivo)
        self.assertTrue(resultado.df.empty)
        self.assertEqual(resultado.linhas_ignoradas, 1)
