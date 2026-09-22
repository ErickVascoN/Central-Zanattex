"""Exportação do relatório de Saldo Fiscal em Excel (openpyxl — já é
dependência da Central, usado em carteira/importador.py)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_CABECALHO_FILL = PatternFill("solid", fgColor="172554")
_CABECALHO_FONTE = Font(color="FFFFFF", bold=True)
_COLUNAS = ["Cliente", "NF", "Produto", "NCM", "Recebido", "Unidade", "Saldo", "% Consumido"]


def gerar_xlsx_saldo(itens: list[dict]) -> bytes:
    """`itens`: mesmo formato usado em relatorio_pdf.gerar_pdf_saldo, com
    `ncm` e `unidade` adicionais."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Saldo Fiscal"

    ws.append(_COLUNAS)
    for celula in ws[1]:
        celula.fill = _CABECALHO_FILL
        celula.font = _CABECALHO_FONTE
        celula.alignment = Alignment(horizontal="center")

    for item in itens:
        ws.append([
            item["cliente"], item["nf"], item["produto"], item.get("ncm", ""),
            float(item["recebido"]), item.get("unidade", ""), float(item["saldo"]),
            round(item["pct"], 1),
        ])

    for coluna, largura in zip("ABCDEFGH", (28, 12, 40, 12, 14, 10, 14, 14)):
        ws.column_dimensions[coluna].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_STATUS_LABEL = {
    "NAO_UTILIZADO": "Não utilizado", "PARCIAL": "Parcialmente utilizado",
    "TOTAL": "Totalmente utilizado", "EXCEDIDO": "Excedido", "DIVERGENCIA": "Com divergência",
}
_COLUNAS_HISTORICO = [
    "NF", "Data", "Fornecedor", "Produto", "Código", "Unidade",
    "Recebido", "Utilizado", "Saldo", "Valor entrada", "Valor utilizado", "Valor saldo", "Status",
]


def gerar_xlsx_historico(itens: list) -> bytes:
    """`itens`: list[servicos.ItemHistorico] — mesma base da tela de
    Histórico, exportada linha a linha."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Histórico"

    ws.append(_COLUNAS_HISTORICO)
    for celula in ws[1]:
        celula.fill = _CABECALHO_FILL
        celula.font = _CABECALHO_FONTE
        celula.alignment = Alignment(horizontal="center")

    for linha in itens:
        item = linha.item
        ws.append([
            item.nota_fiscal.n_nf, item.nota_fiscal.data_emissao.strftime("%d/%m/%Y"),
            item.nota_fiscal.cliente.nome, item.x_prod, item.c_prod, item.u_com,
            float(item.q_com), float(linha.utilizado), float(linha.saldo),
            float(linha.valor_entrada), float(linha.valor_utilizado), float(linha.valor_saldo),
            _STATUS_LABEL.get(linha.status, linha.status),
        ])

    for coluna, largura in zip("ABCDEFGHIJKLM", (12, 12, 26, 40, 16, 9, 12, 12, 12, 14, 14, 14, 20)):
        ws.column_dimensions[coluna].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


